"""Lectura del export de Fanero y cálculo de incentivos por cuota.

Reglas acordadas:
- Vendedor = LOGIN. Líder = LEADER; su venta es la suma de todo su equipo según
  el Fanero (todas las filas donde figura como LEADER, incluidas las suyas).
- Productos:
    PREPAGO       -> columna N, TIPO ACTIVACION = PREPAGO
    POSTPAGO      -> columna N, TIPO ACTIVACION = POSTPAGO
    PORTA PREPAGO -> columna P, PRODUCTO = PORTA PREPAGO
    OSS           -> columna P, PRODUCTO = OSS
  (una porta prepago también cuenta como prepago y un OSS como postpago,
  porque así figuran en TIPO ACTIVACION).
- Con varios productos, la persona debe cumplir la cuota de CADA producto.
- Fecha = LASTSTATUSMODDATE. Rango horario = REQUESTSTARTHOUR (hora de solicitud).
- "Primeros por departamento" = los primeros en LLEGAR a la cuota (a todas las
  cuotas de producto). El momento se toma de LASTSTATUSMODDATE + LASTSTATUSMODHOUR.
- Lista de participantes por incentivo: si existe, solo participan los de la lista.
"""
from __future__ import annotations

import io
import unicodedata

import pandas as pd

COLUMNAS_REQUERIDAS = [
    "STATUS", "REQUESTSTARTHOUR", "LASTSTATUSMODDATE", "LASTSTATUSMODHOUR",
    "LOGIN", "LEADER", "DEPARTMENT", "TIPO ACTIVACION", "PRODUCTO",
]
COLUMNAS_OPCIONALES = [
    "ACTIVATIONTYPE", "TARIFFPLANNAME", "PARTNER", "USERTYPE", "LATITUDE",
    "LONGITUDE", "PORTABILIDAD", "CANAL", "PROVINCIA", "DISTRITO", "FUENTE UBICACION",
]

# producto -> (columna del Fanero, valor)
PRODUCTOS = {
    "PREPAGO": ("TIPO ACTIVACION", "PREPAGO"),
    "POSTPAGO": ("TIPO ACTIVACION", "POSTPAGO"),
    "PORTA PREPAGO": ("PRODUCTO", "PORTA PREPAGO"),
    "OSS": ("PRODUCTO", "OSS"),
}


def slug(producto: str) -> str:
    return producto.replace(" ", "_")


# Nombres de columna en la base (sin espacios)
MAPA_BD = {c: c.lower().replace(" ", "_") for c in COLUMNAS_REQUERIDAS + COLUMNAS_OPCIONALES}


def sin_tildes(texto) -> str:
    if texto is None or (not isinstance(texto, str) and pd.isna(texto)):
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    return "".join(ch for ch in t if not unicodedata.combining(ch)).strip().upper()


def _txt(serie: pd.Series) -> pd.Series:
    """Texto limpio; celdas vacías -> "" (pandas 3 conserva NaN al hacer astype(str))."""
    s = serie.astype(object)
    return s.where(s.notna(), "").map(str).str.strip().replace({"nan": "", "None": "", "<NA>": "", "NaT": ""})


def _limpiar_id(serie: pd.Series) -> pd.Series:
    return _txt(serie).str.replace(r"\.0$", "", regex=True).str.lstrip("'").str.strip()


def _decodificar(contenido: bytes) -> str:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return contenido.decode(enc)
        except UnicodeDecodeError:
            continue
    return contenido.decode("utf-8", errors="ignore")


def leer_tabla(nombre: str, contenido: bytes, hoja=0) -> pd.DataFrame:
    """Lee xlsx, xls, csv o txt separado por tabs/;/, todo como texto."""
    nombre = nombre.lower()
    if nombre.endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(io.BytesIO(contenido), dtype=str, sheet_name=hoja)
    texto = _decodificar(contenido)
    primera = texto.splitlines()[0] if texto else ""
    sep = "\t" if "\t" in primera else (";" if primera.count(";") > primera.count(",") else ",")
    return pd.read_csv(io.StringIO(texto), sep=sep, dtype=str)


leer_archivo = leer_tabla


def normalizar(df: pd.DataFrame) -> pd.DataFrame:
    """Valida columnas y deja el export de Fanero listo para guardar/calcular."""
    df = df.copy()
    df.columns = [str(c).strip().upper() for c in df.columns]
    faltan = [c for c in COLUMNAS_REQUERIDAS if c not in df.columns]
    if faltan:
        raise ValueError("Al archivo le faltan estas columnas: " + ", ".join(faltan))

    for c in COLUMNAS_OPCIONALES:
        if c not in df.columns:
            df[c] = None
    df = df[COLUMNAS_REQUERIDAS + COLUMNAS_OPCIONALES]
    df = df.dropna(how="all")

    df["LOGIN"] = _limpiar_id(df["LOGIN"])
    df["LEADER"] = _limpiar_id(df["LEADER"])
    for c in ("STATUS", "PROVINCIA", "DISTRITO", "CANAL", "USERTYPE", "PORTABILIDAD",
              "ACTIVATIONTYPE", "TARIFFPLANNAME", "PARTNER", "FUENTE UBICACION"):
        df[c] = _txt(df[c]).replace({"": None})
    for c in ("TIPO ACTIVACION", "PRODUCTO", "DEPARTMENT"):
        df[c] = df[c].map(sin_tildes)
    for c in ("LATITUDE", "LONGITUDE"):
        df[c] = pd.to_numeric(_txt(df[c]).str.replace("'", "", regex=False), errors="coerce")

    df["REQUESTSTARTHOUR"] = pd.to_numeric(df["REQUESTSTARTHOUR"], errors="coerce").astype("Int64")
    fecha = pd.to_datetime(_txt(df["LASTSTATUSMODDATE"]).str[:10],
                           dayfirst=True, errors="coerce")
    hora = _txt(df["LASTSTATUSMODHOUR"])
    df["LASTSTATUSMODDATE"] = fecha.dt.date
    df["FECHA_HORA"] = pd.to_datetime(fecha.dt.strftime("%Y-%m-%d") + " " + hora, errors="coerce")
    df["FECHA_HORA"] = df["FECHA_HORA"].fillna(fecha)

    malas = df["LASTSTATUSMODDATE"].isna() | (df["LOGIN"] == "")
    return df[~malas].reset_index(drop=True)


def a_formato_bd(df: pd.DataFrame) -> pd.DataFrame:
    out = df.rename(columns=MAPA_BD).rename(columns={"FECHA_HORA": "fecha_hora"})
    return out.rename(columns={"laststatusmoddate": "fecha"})


def desde_bd(df: pd.DataFrame) -> pd.DataFrame:
    inv = {v: k for k, v in MAPA_BD.items()}
    inv.update({"fecha": "LASTSTATUSMODDATE", "fecha_hora": "FECHA_HORA"})
    out = df.rename(columns=inv)
    out["FECHA_HORA"] = pd.to_datetime(out["FECHA_HORA"])
    out["LASTSTATUSMODDATE"] = pd.to_datetime(out["LASTSTATUSMODDATE"]).dt.date
    out["REQUESTSTARTHOUR"] = pd.to_numeric(out["REQUESTSTARTHOUR"], errors="coerce").astype("Int64")
    for c in ("TIPO ACTIVACION", "PRODUCTO", "DEPARTMENT"):
        out[c] = out[c].map(sin_tildes)
    return out


# --------------------------------------------------------------------------- #
# Participantes
# --------------------------------------------------------------------------- #
COLS_PARTICIPANTES = (["DNI", "NOMBRE", "ROL", "DEPARTAMENTO", "DNI_LIDER"]
                      + [f"CUOTA_{slug(p)}" for p in PRODUCTOS])


def plantilla_participantes() -> bytes:
    ejemplo = pd.DataFrame([
        ["40854375", "Juan Pérez", "Vendedor", "LORETO", "41544805", 5, "", 2, ""],
        ["62300640", "María Ríos", "Vendedor", "LORETO", "41544805", "", "", "", ""],
        ["41544805", "Carlos Soto", "Líder", "LORETO", "", 30, "", 10, ""],
    ], columns=COLS_PARTICIPANTES)
    instrucciones = pd.DataFrame([
        ["DNI", "Obligatorio", "El número que sale en LOGIN (vendedor) o LEADER (líder) del Fanero"],
        ["NOMBRE", "Opcional", "Solo para identificarlo en la matriz y el Excel. Puede ir vacío"],
        ["ROL", "Obligatorio", "Vendedor o Líder. Si alguien compite en ambos, va en dos filas"],
        ["DEPARTAMENTO", "Opcional", "Departamento en el que compite. Vacío = donde más vende según el Fanero"],
        ["DNI_LIDER", "Opcional", "Solo vendedores: DNI de su líder. Vacío = el que figura en el Fanero"],
        ["CUOTA_PREPAGO", "Opcional", "Cuota individual de prepago (TIPO ACTIVACION = PREPAGO). Vacío = cuota general del incentivo"],
        ["CUOTA_POSTPAGO", "Opcional", "Cuota individual de postpago (TIPO ACTIVACION = POSTPAGO). Vacío = cuota general"],
        ["CUOTA_PORTA_PREPAGO", "Opcional", "Cuota individual de porta prepago (PRODUCTO = PORTA PREPAGO). Vacío = cuota general"],
        ["CUOTA_OSS", "Opcional", "Cuota individual de OSS (PRODUCTO = OSS). Vacío = cuota general"],
        ["Nota", "", "Solo se usan las cuotas de los productos que tenga el incentivo. Una cuota 0 = ese producto no se le exige"],
    ], columns=["Columna", "¿Obligatoria?", "Qué poner"])
    return _excel({"Participantes": ejemplo, "Instrucciones": instrucciones}, texto_cols={"A", "E"})


def leer_participantes(nombre: str, contenido: bytes) -> tuple[list[dict], list[str]]:
    """Devuelve (participantes, avisos). Lanza ValueError si el archivo no sirve."""
    df = leer_tabla(nombre, contenido, hoja=0)
    df.columns = [sin_tildes(c).replace(" ", "_") for c in df.columns]
    faltan = [c for c in ("DNI", "ROL") if c not in df.columns]
    if faltan:
        raise ValueError("Al archivo de participantes le faltan: " + ", ".join(faltan))
    for c in COLS_PARTICIPANTES:
        if c not in df.columns:
            df[c] = None
    df = df[COLS_PARTICIPANTES].dropna(how="all")
    avisos = []

    df["DNI"] = _limpiar_id(df["DNI"])
    df["DNI_LIDER"] = _limpiar_id(df["DNI_LIDER"])
    df["NOMBRE"] = _txt(df["NOMBRE"])
    df["DEPARTAMENTO"] = df["DEPARTAMENTO"].map(sin_tildes)
    rol = df["ROL"].map(sin_tildes)
    df["ROL"] = rol.map(lambda r: "Vendedor" if r.startswith("VEND") else ("Líder" if r.startswith("LID") else None))

    # Líder con DNI vacío pero con su número en DNI_LIDER: se toma de ahí
    es_lider = df["ROL"] == "Líder"
    mover = es_lider & (df["DNI"] == "") & (df["DNI_LIDER"] != "")
    if mover.any():
        df.loc[mover, "DNI"] = df.loc[mover, "DNI_LIDER"]
    df.loc[es_lider, "DNI_LIDER"] = ""

    sin_dni = df["DNI"] == ""
    if sin_dni.any():
        avisos.append(f"{int(sin_dni.sum())} fila(s) sin DNI se ignoraron.")
    malos = ~sin_dni & df["ROL"].isna()
    if malos.any():
        raise ValueError("ROL debe ser Vendedor o Líder. Revisa los DNI: "
                         + ", ".join(map(str, df.loc[malos, "DNI"].head(10))))
    df = df[~sin_dni]

    for p in PRODUCTOS:
        c = f"CUOTA_{slug(p)}"
        txt = _txt(df[c]).str.replace(r"\.0$", "", regex=True)
        num = pd.to_numeric(txt, errors="coerce")
        raros = (txt != "") & num.isna()
        if raros.any():
            raise ValueError(f"{c} tiene valores que no son números (DNI {', '.join(df.loc[raros, 'DNI'].head(5))}).")
        df[c] = num.map(lambda x: None if pd.isna(x) else int(x))

    dup = df.duplicated(["DNI", "ROL"], keep="first")
    if dup.any():
        avisos.append(f"{int(dup.sum())} DNI repetido(s) con el mismo rol; se tomó la primera fila: "
                      + ", ".join(map(str, df.loc[dup, "DNI"].head(10))))
        df = df[~dup]
    registros = df.astype(object).where(df.notna(), None).to_dict("records")
    return registros, avisos


# --------------------------------------------------------------------------- #
# Cálculo
# --------------------------------------------------------------------------- #
CONFIG_DEFECTO = {
    "nombre": "",
    "fecha_inicio": None,          # "YYYY-MM-DD"
    "fecha_fin": None,
    "hora_desde": 0,               # REQUESTSTARTHOUR inclusive
    "hora_hasta": 23,              # inclusive
    "evaluacion": "periodo",       # "periodo" o "diaria"
    "estados": ["Completada"],
    "departamentos": [],           # vacío = todos
    "productos": ["PREPAGO"],      # subconjunto de PRODUCTOS
    "participa_vendedor": True,
    "participa_lider": True,
    "cuotas_vendedor": {"PREPAGO": 1},
    "cuotas_lider": {"PREPAGO": 1},
    "gana_todos_cumplen": True,
    "gana_primeros_depto": False,
    "primeros_n": 1,
    "premio_vendedor": "",
    "premio_lider": "",
    "participantes": [],           # lista por incentivo; vacía = todos los del Fanero
}


def completar_config(cfg: dict) -> dict:
    cfg = {**CONFIG_DEFECTO, **(cfg or {})}
    cfg["productos"] = [p for p in cfg.get("productos") or [] if p in PRODUCTOS] or ["PREPAGO"]
    for clave in ("cuotas_vendedor", "cuotas_lider"):
        cfg[clave] = {p: int(cfg.get(clave, {}).get(p, 1) or 0) for p in cfg["productos"]}
    return cfg


def mascara_producto(df: pd.DataFrame, producto: str) -> pd.Series:
    col, valor = PRODUCTOS[producto]
    return df[col] == valor


def filtrar(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Ventas que entran al incentivo (fechas, horas, estado, departamento).
    El filtro de producto se aplica por producto en el cálculo."""
    d = df
    if cfg.get("fecha_inicio"):
        d = d[d["LASTSTATUSMODDATE"] >= pd.to_datetime(cfg["fecha_inicio"]).date()]
    if cfg.get("fecha_fin"):
        d = d[d["LASTSTATUSMODDATE"] <= pd.to_datetime(cfg["fecha_fin"]).date()]
    h1, h2 = int(cfg.get("hora_desde", 0)), int(cfg.get("hora_hasta", 23))
    d = d[d["REQUESTSTARTHOUR"].between(h1, h2)]
    if cfg.get("estados"):
        d = d[d["STATUS"].isin(cfg["estados"])]
    if cfg.get("departamentos"):
        d = d[d["DEPARTMENT"].isin(cfg["departamentos"])]
    return d


def _moda(serie: pd.Series) -> str:
    s = serie[_txt(serie) != ""]
    m = s.mode()
    return m.iloc[0] if not m.empty else ""


def _personas(universo: pd.DataFrame, rol: str, cfg: dict) -> pd.DataFrame:
    """Tabla de participantes de un rol con DNI, NOMBRE, DEPARTAMENTO, LIDER y Q_<producto>."""
    col = "LOGIN" if rol == "Vendedor" else "LEADER"
    general = cfg["cuotas_vendedor"] if rol == "Vendedor" else cfg["cuotas_lider"]
    info = (universo[universo[col] != ""].groupby(col)
            .agg(DEPTO_FANERO=("DEPARTMENT", _moda), LIDER_FANERO=("LEADER", _moda))
            .reset_index().rename(columns={col: "DNI"}))

    lista = [p for p in cfg.get("participantes") or [] if p.get("ROL") == rol]
    if cfg.get("participantes"):
        base = pd.DataFrame(lista, columns=COLS_PARTICIPANTES)
        base = base.merge(info, on="DNI", how="left")
        base["DEPARTAMENTO"] = base["DEPARTAMENTO"].where(base["DEPARTAMENTO"].fillna("") != "",
                                                          base["DEPTO_FANERO"]).fillna("")
        base["LIDER"] = base["DNI_LIDER"].where(base["DNI_LIDER"].fillna("") != "",
                                                base["LIDER_FANERO"]).fillna("")
        for p in cfg["productos"]:
            ind = pd.to_numeric(base[f"CUOTA_{slug(p)}"], errors="coerce")
            base[f"Q_{p}"] = ind.fillna(general[p]).astype(int)
    else:
        base = info.rename(columns={"DEPTO_FANERO": "DEPARTAMENTO", "LIDER_FANERO": "LIDER"})
        base["NOMBRE"] = ""
        for p in cfg["productos"]:
            base[f"Q_{p}"] = int(general[p])
    if rol == "Líder":
        base["LIDER"] = ""
    return base[["DNI", "NOMBRE", "DEPARTAMENTO", "LIDER"] + [f"Q_{p}" for p in cfg["productos"]]]


def _matriz_rol(d: pd.DataFrame, universo: pd.DataFrame, rol: str, cfg: dict) -> pd.DataFrame:
    col = "LOGIN" if rol == "Vendedor" else "LEADER"
    diaria = cfg.get("evaluacion") == "diaria"
    claves = ["DNI"] + (["DIA"] if diaria else [])
    prods = cfg["productos"]

    m = _personas(universo, rol, cfg)
    if diaria:
        if cfg.get("fecha_inicio") and cfg.get("fecha_fin"):
            dias = list(pd.date_range(cfg["fecha_inicio"], cfg["fecha_fin"]).date)
        else:
            dias = sorted(d["LASTSTATUSMODDATE"].unique())
        m = m.merge(pd.DataFrame({"DIA": dias}), how="cross")

    d = d[d[col] != ""].sort_values("FECHA_HORA").copy()
    d["DNI"] = d[col]
    d["DIA"] = d["LASTSTATUSMODDATE"]
    for p in prods:
        dp = d[mascara_producto(d, p)].copy()
        conteo = dp.groupby(claves).size().rename(f"V_{p}").reset_index()
        m = m.merge(conteo, on=claves, how="left")
        m[f"V_{p}"] = m[f"V_{p}"].fillna(0).astype(int)
        # momento en que llega a SU cuota de este producto
        dp["_N"] = dp.groupby(claves).cumcount() + 1
        dp = dp.merge(m[claves + [f"Q_{p}"]], on=claves, how="inner")
        llega = (dp[dp["_N"] == dp[f"Q_{p}"]][claves + ["FECHA_HORA"]]
                 .rename(columns={"FECHA_HORA": f"T_{p}"}))
        m = m.merge(llega, on=claves, how="left")

    q = m[[f"Q_{p}" for p in prods]].to_numpy()
    v = m[[f"V_{p}" for p in prods]].to_numpy()
    m["TOTAL_VENTAS"] = v.sum(axis=1)
    cumple_p = v >= q
    m["CUMPLE"] = cumple_p.all(axis=1)
    pct = pd.DataFrame(
        {p: (m[f"V_{p}"] / m[f"Q_{p}"].where(m[f"Q_{p}"] > 0) * 100) for p in prods})
    m["AVANCE_%"] = pct.clip(upper=100).min(axis=1, skipna=True).fillna(100).round(1)
    m["FALTAN"] = (q - v).clip(min=0).sum(axis=1)
    tiempos = m[[f"T_{p}" for p in prods]].where(m[[f"Q_{p}" for p in prods]].gt(0).to_numpy())
    m["HORA_CUOTA"] = tiempos.max(axis=1).where(m["CUMPLE"])

    grupo = ["DEPARTAMENTO"] + (["DIA"] if diaria else [])
    m["PUESTO_DEPTO"] = (m[m["CUMPLE"]].groupby(grupo)["HORA_CUOTA"]
                         .rank(method="first").reindex(m.index))
    gana_todos = m["CUMPLE"] & bool(cfg.get("gana_todos_cumplen"))
    gana_primero = m["PUESTO_DEPTO"].le(int(cfg.get("primeros_n", 1))) & bool(cfg.get("gana_primeros_depto"))
    m["GANA"] = gana_todos | gana_primero
    m["MOTIVO"] = [" + ".join(r for r, ok in (("Cumplió cuota", t), ("Primero(s) del departamento", p)) if ok)
                   for t, p in zip(gana_todos, gana_primero)]
    premio = cfg.get("premio_vendedor" if rol == "Vendedor" else "premio_lider", "")
    m["PREMIO"] = m["GANA"].map(lambda g: premio if g else "")
    m.insert(0, "ROL", rol)
    return m


def calcular(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    cfg = completar_config(cfg)
    universo = df
    if cfg.get("estados"):
        universo = universo[universo["STATUS"].isin(cfg["estados"])]
    d = filtrar(df, cfg)

    partes = []
    if cfg["participa_vendedor"]:
        partes.append(_matriz_rol(d, universo, "Vendedor", cfg))
    if cfg["participa_lider"]:
        partes.append(_matriz_rol(d, universo, "Líder", cfg))
    if not partes:
        return pd.DataFrame()
    m = pd.concat(partes, ignore_index=True)

    diaria = cfg["evaluacion"] == "diaria"
    renombre, cols = {}, ["ROL", "DNI", "NOMBRE", "DEPARTAMENTO", "LIDER"] + (["DIA"] if diaria else [])
    for p in cfg["productos"]:
        renombre[f"V_{p}"] = f"{p} VENTAS"
        renombre[f"Q_{p}"] = f"{p} CUOTA"
        cols += [f"V_{p}", f"Q_{p}"]
    cols += ["TOTAL_VENTAS", "AVANCE_%", "FALTAN", "CUMPLE", "HORA_CUOTA", "PUESTO_DEPTO",
             "GANA", "MOTIVO", "PREMIO"]
    m = m[cols].rename(columns=renombre)
    m["PUESTO_DEPTO"] = m["PUESTO_DEPTO"].astype("Int64")
    m["CUMPLE"] = m["CUMPLE"].map({True: "Sí", False: "No"})
    m["GANA"] = m["GANA"].map({True: "Sí", False: "No"})
    orden = ["ROL", "DEPARTAMENTO"] + (["DIA"] if diaria else [])
    m = m.sort_values(orden + ["GANA", "PUESTO_DEPTO", "AVANCE_%"],
                      ascending=[True] * len(orden) + [False, True, False], na_position="last")
    return m.reset_index(drop=True)


def resumen(m: pd.DataFrame) -> pd.DataFrame:
    if m.empty:
        return m
    return (m.groupby(["ROL", "DEPARTAMENTO"])
            .agg(PARTICIPANTES=("DNI", "nunique"),
                 VENTAS=("TOTAL_VENTAS", "sum"),
                 CUMPLEN=("CUMPLE", lambda s: (s == "Sí").sum()),
                 GANADORES=("GANA", lambda s: (s == "Sí").sum()))
            .reset_index())


def reglas_tabla(cfg: dict) -> pd.DataFrame:
    cfg = completar_config(cfg)
    filas = []
    for k, v in cfg.items():
        if k == "participantes":
            v = f"{len(v)} en lista" if v else "Todos los del Fanero"
        elif isinstance(v, dict):
            v = ", ".join(f"{p}: {n}" for p, n in v.items())
        elif isinstance(v, list):
            v = ", ".join(map(str, v))
        filas.append((k, v))
    filas.append(("regla_productos", "PREPAGO/POSTPAGO por TIPO ACTIVACION; PORTA PREPAGO/OSS por PRODUCTO"))
    return pd.DataFrame(filas, columns=["Parámetro", "Valor"])


def _excel(hojas: dict, texto_cols=frozenset()) -> bytes:
    from openpyxl.styles import Font, PatternFill
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre, index=False)
        azul = PatternFill("solid", fgColor="00679E")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = azul
            for col in ws.columns:
                ancho = max(len(str(c.value or "")) for c in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, ancho + 2), 60)
                if col[0].column_letter in texto_cols:
                    for c in col[1:]:
                        c.number_format = "@"
            ws.freeze_panes = "A2"
    return buf.getvalue()


def a_excel(matriz: pd.DataFrame, cfg: dict) -> bytes:
    m = matriz.copy()
    if "HORA_CUOTA" in m:
        m["HORA_CUOTA"] = pd.to_datetime(m["HORA_CUOTA"]).dt.strftime("%d/%m/%Y %H:%M:%S").fillna("")
    return _excel({"Matriz": m, "Ganadores": m[m["GANA"] == "Sí"],
                   "Resumen depto": resumen(matriz), "Reglas": reglas_tabla(cfg)})
